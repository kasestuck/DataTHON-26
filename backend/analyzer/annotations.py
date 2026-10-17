"""Preserve human spreadsheet observations; apply locally proposed visual groups.

source_observed records grouping of a human observation, not human approval of
an AI proposal and never a legal finding. Ambiguity stays outside typed scoring.
"""
from collections import Counter
import hashlib
import json
from pathlib import Path
import re

from .catalog import validate_catalog
from .labels import import_open_labels, parse_timecode, parse_verdict, video_key

UNCERTAIN = re.compile(
    r"предполож|возможн|вероят|кажется|не\s*понят|не\s+видно|неясн|подозр|может\s+быть|вроде|якобы|\?",
    re.IGNORECASE)


def apply_groups(source, videos_dir, proposals, *, description_column="D"):
    """Apply only groups explicitly supported by non-speculative human text.

    Does not silently accept source changes, duplicate rows, invented assignments,
    missing observations, unknown IDs or negative-to-positive contradictions.
    """
    from openpyxl import load_workbook

    digest = hashlib.sha256(Path(source).read_bytes()).hexdigest()
    if proposals.get("source_sha256") != digest:
        raise ValueError("Grouping proposals do not match the source workbook")
    catalog = validate_catalog(proposals["catalog"])
    labels, report = import_open_labels(source, videos_dir)
    book = load_workbook(source, data_only=False)
    try:
        sheet = book.active
        source_rows = {video_key(sheet.cell(i, 1).value): i for i in range(2, sheet.max_row + 1)
                       if sheet.cell(i, 1).value is not None}
        assignments = proposals["assignments"]
        by_row = {value["source_row"]: value for value in assignments}
        if len(by_row) != len(assignments) or set(by_row) != set(source_rows.values()):
            raise ValueError("Grouping must cover each source row exactly once")
        for row in labels:
            number = source_rows[video_key(row["video"])]
            cell = sheet[f"{description_column}{number}"]
            text = cell.value if cell.value is not None else ""
            if not isinstance(text, str):
                raise ValueError("Observations must be text; preserve other cell types in the source archive")
            assignment = by_row[number]
            kind, status = assignment["type_id"], assignment["status"]
            if status not in {"source_observed", "needs_review", "negative"}:
                raise ValueError("Invalid annotation status")
            if kind is not None and kind not in catalog:
                raise ValueError("Proposed type is not in the catalog")
            positive = parse_verdict(row["has_violation"])
            reasons = []
            if not positive and (kind is not None or status == "source_observed"):
                reasons.append("grouping_contradicts_negative_source")
            if UNCERTAIN.search(text):
                reasons.append("speculative_or_unclear_source_description")
            if cell.data_type == "f":
                reasons.append("description_is_a_formula_review_original")
            if positive and (not text.strip() or parse_timecode(row["timecode_sec"]) is None):
                reasons.append("missing_observation_or_timestamp")
            if positive and (kind is None or status != "source_observed"):
                reasons.append("visual_type_not_established")
            if status == "needs_review":
                reasons.append("local_grouping_requests_human_review")
            row.update(observed_action=text, source_row=number, source_sha256=digest,
                       proposed_type=kind or "", notes=";".join(dict.fromkeys(reasons)))
            if reasons:
                row.update(review_status="needs_review", violation_type="unknown" if positive else "")
            elif positive:
                row.update(review_status="source_observed", violation_type=kind)
            else:
                row.update(review_status="negative", violation_type="", timecode_sec="")
    finally:
        book.close()
    counts = Counter(row["violation_type"] for row in labels if row["review_status"] == "source_observed")
    # A proposed category without any explicit human observation stays inactive.
    active_catalog = {kind: item for kind, item in catalog.items() if counts[kind]}
    report.update(source_sha256=digest, examples_per_type=dict(counts),
        needs_review=[{"video": row["video"], "timecode_sec": row["timecode_sec"],
                       "source_row": row["source_row"], "reasons": row["notes"].split(";")}
                      for row in labels if row["review_status"] == "needs_review"],
        label_provenance="Human observations, locally assisted taxonomy; no legal confirmation",
        explicit_human_type_approval=False)
    return labels, active_catalog, report


def load_proposals(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def confirm_groups(labels, catalog, confirmation=None):
    """Only an explicit human review promotes proposed IDs into ground truth.

    The grouping model never writes this confirmation document. Its source hash
    binds human decisions to the exact workbook reviewed, including row numbers.
    """
    confirmed = {}
    if confirmation is not None:
        digests = {row["source_sha256"] for row in labels}
        if digests != {confirmation.get("source_sha256")}:
            raise ValueError("Human review does not match the source workbook")
        by_row = {int(row["source_row"]): row for row in labels}
        for event in confirmation["events"]:
            number, kind = event["source_row"], event["type_id"]
            if (type(number) is not int or number in confirmed or number not in by_row
                    or not parse_verdict(by_row[number]["has_violation"]) or kind not in catalog):
                raise ValueError("Invalid or duplicate human type confirmation")
            confirmed[number] = kind
    result = []
    for original in labels:
        row = dict(original)
        number = int(row["source_row"])
        if number in confirmed:
            row.update(violation_type=confirmed[number], review_status="confirmed")
        elif parse_verdict(row["has_violation"]):
            if row["review_status"] == "source_observed":
                row["notes"] = ";".join(filter(None, [row["notes"], "manual_type_approval_pending"]))
            row.update(violation_type="unknown", review_status="needs_review")
        result.append(row)
    counts = Counter(confirmed.values())
    return result, {kind: item for kind, item in catalog.items() if counts[kind]}
