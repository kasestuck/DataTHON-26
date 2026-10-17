"""Apply explicit user decisions, without converting preliminary labels to truth."""
import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path

from backend.analyzer.catalog import load_catalog
from backend.analyzer.labels import import_open_labels, parse_verdict, video_key, write_labels
from backend.scripts.prepare_annotations import snapshot_workbook


def finalize(source, videos_dir, decisions, catalog):
    from openpyxl import load_workbook
    digest = hashlib.sha256(Path(source).read_bytes()).hexdigest()
    if decisions["source_sha256"] != digest:
        raise ValueError("Manual decisions refer to a different workbook")
    labels, report = import_open_labels(source, videos_dir)
    choices = {video_key(key): value for key, value in decisions["decisions"].items()}
    if len(choices) != len(decisions["decisions"]):
        raise ValueError("Duplicate normalized decision names")
    positives = {video_key(r["video"]) for r in labels if parse_verdict(r["has_violation"])}
    if set(choices) != positives:
        raise ValueError("Manual decisions must cover every positive video, and only positives")
    book = load_workbook(source, data_only=False)
    try:
        sheet = book.active
        original = {video_key(sheet.cell(i, 1).value): (i, sheet.cell(i, 4).value or "")
                    for i in range(2, sheet.max_row + 1) if sheet.cell(i, 1).value}
        for row in labels:
            key = video_key(row["video"])
            number, text = original[key]
            row.update(source_observed_action=text, observed_action=text, source_row=number,
                source_sha256=digest, organizer_has_violation=row["has_violation"],
                organizer_timecode_sec=row["timecode_sec"], legal_basis_status="not_applicable",
                review_status="negative", annotation_source="organizer_workbook", proposed_type="")
            if key not in choices:
                continue
            value = choices[key]
            kind, status = value["type_id"], value["review_status"]
            if kind not in catalog or status not in {"confirmed_visual", "needs_review"}:
                raise ValueError("Invalid manual type or review status")
            row.update(proposed_type=kind, review_status=status, legal_basis_status="unverified",
                annotation_source=value.get("annotation_source", "user_preliminary_mapping"),
                notes=value.get("notes", ""), observed_action=value.get("observed_action", text),
                violation_type=kind if status == "confirmed_visual" else "unknown")
    finally:
        book.close()
    if len(labels) != decisions["expected_rows"] or len(positives) != decisions["expected_positive"]:
        raise ValueError("Dataset coverage changed; review the new source")
    report.update(confirmed_visual=sum(r["review_status"] == "confirmed_visual" for r in labels),
        needs_review=sum(r["review_status"] == "needs_review" for r in labels),
        proposed_per_type=dict(Counter(r["proposed_type"] for r in labels if r["proposed_type"])),
        confirmed_per_type=dict(Counter(r["violation_type"] for r in labels if r["review_status"] == "confirmed_visual")))
    return labels, report


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--excel", default="dataset/open_labels.xlsx")
    parser.add_argument("--videos-dir", default="dataset/videos")
    parser.add_argument("--decisions", default="backend/data/annotations/final_review.json")
    parser.add_argument("--catalog", default="backend/data/violation_types.json")
    parser.add_argument("--output", default="backend/data/labels.csv")
    args = parser.parse_args()
    archive = Path(args.decisions).parent
    snapshot_workbook(args.excel, archive)
    output = Path(args.output)
    if output.exists():
        data = output.read_bytes()
        saved = archive / ("labels_before_" + hashlib.sha256(data).hexdigest() + ".csv")
        if not saved.exists():
            saved.write_bytes(data)
    labels, report = finalize(args.excel, args.videos_dir,
        json.loads(Path(args.decisions).read_text(encoding="utf-8")), load_catalog(args.catalog))
    write_labels(labels, output)
    write_labels([r for r in labels if r["review_status"] == "needs_review"], output.with_name("needs_review.csv"))
    reports = output.parent / "reports"
    reports.mkdir(exist_ok=True)
    (reports / "annotation_finalization.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=True))


if __name__ == "__main__":
    main()
